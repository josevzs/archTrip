"""Excel templates (download) and parsing of the filled-in files (upload)."""
import io
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# (header shown in the template, internal field, required, width, example value)
ROUTE_COLUMNS = [
    ("Orden", "position", True, 8, 1),
    ("Ciudad", "city", True, 24, "Oporto"),
    ("País", "country", False, 16, "Portugal"),
    ("Notas", "notes", False, 40, "Noche 1 y 2"),
]

LANDMARK_COLUMNS = [
    ("Edificio", "name", True, 32, "Casa da Música"),
    ("Arquitecto", "architect", True, 28, "Rem Koolhaas / OMA"),
    ("Ciudad", "city", True, 20, "Oporto"),
    ("Dirección", "address", False, 32, "Av. da Boavista 604"),
    ("Año", "year", False, 8, 2005),
    ("Latitud", "lat", False, 12, ""),
    ("Longitud", "lon", False, 12, ""),
    ("URL ArchDaily", "url_archdaily", False, 36, ""),
    ("URL Arquitectura Viva", "url_av", False, 36, ""),
    ("URL Imagen 1", "url_image1", False, 36, ""),
    ("URL Imagen 2", "url_image2", False, 36, ""),
    ("Notas", "notes", False, 40, "A · ejemplo de nota"),
    ("Estado", "status", False, 12, ""),
]
STATUSES = ("pendiente", "curado", "posible", "descartado")

# Accepted header spellings per field (normalised: lowercase, no accents).
ROUTE_ALIASES = {
    "position": {"orden", "n", "num", "numero", "posicion", "dia"},
    "city": {"ciudad", "localidad", "pueblo", "ciudad/pueblo", "lugar"},
    "country": {"pais"},
    "notes": {"notas", "nota", "comentarios", "observaciones"},
}

LANDMARK_ALIASES = {
    "name": {"edificio", "nombre", "nombre del edificio", "obra", "proyecto", "hito", "landmark"},
    "architect": {"arquitecto", "arquitectos", "arquitecto/a", "autor", "estudio"},
    "city": {"ciudad", "localidad", "pueblo", "ciudad/localidad", "lugar", "municipio"},
    "address": {"direccion", "calle"},
    "year": {"ano", "anio", "year", "fecha"},
    "lat": {"latitud", "lat"},
    "lon": {"longitud", "lon", "lng", "long"},
    "url_archdaily": {"url archdaily", "archdaily", "enlace archdaily"},
    "url_av": {"url arquitectura viva", "arquitectura viva", "av", "url av", "enlace arquitectura viva"},
    "url_image1": {"url imagen 1", "imagen 1", "imagen", "url imagen", "foto", "foto 1", "url foto"},
    "url_image2": {"url imagen 2", "imagen 2", "foto 2", "url foto 2"},
    "notes": {"notas", "nota", "comentarios", "observaciones"},
    "status": {"estado", "status", "curado"},
}

ROUTE_INSTRUCTIONS = [
    ("Orden", "Número de parada en el viaje (1, 2, 3…). Obligatorio."),
    ("Ciudad", "Ciudad o pueblo donde se hace noche o parada. Obligatorio."),
    ("País", "Ayuda a localizar la ciudad en el mapa. Recomendado."),
    ("Notas", "Lo que quieras: fechas, hotel, etc."),
]

LANDMARK_INSTRUCTIONS = [
    ("Edificio", "Nombre del edificio u obra. Obligatorio."),
    ("Arquitecto", "Autor o estudio. Obligatorio."),
    ("Ciudad", "Ciudad o pueblo donde está. Obligatorio; se usa para localizarlo."),
    ("Dirección", "Calle y número si la sabes: mejora mucho la localización automática."),
    ("Año", "Año de construcción (solo informativo)."),
    ("Latitud / Longitud", "Opcionales. Si las rellenas, no se busca la ubicación automáticamente."),
    ("URL ArchDaily", "Enlace a la ficha en archdaily.com. Si se deja vacío, la herramienta la busca."),
    ("URL Arquitectura Viva", "Enlace a arquitecturaviva.com. Si se deja vacío, la herramienta la busca."),
    ("URL Imagen 1 / 2", "Enlaces directos a imágenes (jpg, png…). Si se dejan vacíos, se ofrece una búsqueda de imágenes."),
    ("Notas", "Empieza por la prioridad y los avisos, p. ej. \"A [R] · motivo\". A: justifica desviar el viaje; B: muy recomendable en la zona; C: para completistas. [R] reserva/calendario, [E] solo exterior, [X] no visitable."),
    ("Estado", "Opcional: posible o descartado (también curado). Solo se aplica a hitos nuevos o todavía pendientes; nunca pisa el curado hecho en la herramienta."),
]


def normalise(text):
    """lowercase, strip accents, collapse whitespace — for headers and match keys."""
    if text is None:
        return ""
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().lower()


def landmark_key(name, architect):
    return f"{normalise(name)}|{normalise(architect)}"


# ---------------------------------------------------------------- templates

def _build_template(columns, instructions, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    header_font = Font(name="Courier New", bold=True)
    header_fill = PatternFill("solid", fgColor="EFEFEF")
    for idx, (header, _field, required, width, example) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header + (" *" if required else ""))
        cell.font = header_font
        cell.fill = header_fill
        ws.cell(row=2, column=idx, value=example if example != "" else None)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Instrucciones")
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 90
    ws2.cell(row=1, column=1, value="Columna").font = header_font
    ws2.cell(row=1, column=2, value="Qué poner").font = header_font
    for r, (col, text) in enumerate(instructions, start=2):
        ws2.cell(row=r, column=1, value=col)
        ws2.cell(row=r, column=2, value=text)
    note_row = len(instructions) + 3
    ws2.cell(row=note_row, column=1, value="Nota")
    ws2.cell(
        row=note_row, column=2,
        value="La fila 2 es un ejemplo: bórrala o sobrescríbela. Las columnas con * son obligatorias. "
              "No cambies los nombres de las cabeceras.",
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def route_template():
    return _build_template(ROUTE_COLUMNS, ROUTE_INSTRUCTIONS, "Ruta")


def landmarks_template():
    return _build_template(LANDMARK_COLUMNS, LANDMARK_INSTRUCTIONS, "Hitos")


# ------------------------------------------------------------------ parsing

def _map_headers(header_cells, aliases):
    """Return {column_index: field} for every recognised header."""
    mapping = {}
    for idx, raw in enumerate(header_cells):
        h = normalise(raw).rstrip("*").strip()
        if not h:
            continue
        for field, names in aliases.items():
            if h in names and field not in mapping.values():
                mapping[idx] = field
                break
    return mapping


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _cell_float(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _read_rows(file_obj):
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    # first row with any content is the header
    start = next((i for i, r in enumerate(all_rows) if any(_cell_str(c) for c in r)), None)
    if start is None:
        return None, []
    return all_rows[start], [(start + 2 + i, r) for i, r in enumerate(all_rows[start + 1:])]


def parse_route(file_obj):
    """-> (stops, errors). stops: [{position, city, country, notes}] ordered."""
    header, body = _read_rows(file_obj)
    if header is None:
        return [], ["El archivo está vacío."]
    mapping = _map_headers(header, ROUTE_ALIASES)
    if "city" not in mapping.values():
        return [], ["No encuentro la columna 'Ciudad'. Usa la plantilla descargada."]

    stops, errors = [], []
    for rownum, cells in body:
        rec = {mapping[i]: cells[i] if i < len(cells) else None for i in mapping}
        if not any(_cell_str(v) for v in rec.values()):
            continue
        city = _cell_str(rec.get("city"))
        if not city:
            errors.append(f"fila {rownum}: falta Ciudad")
            continue
        pos = _cell_float(rec.get("position"))
        stops.append({
            "position": int(pos) if pos is not None else None,
            "city": city,
            "country": _cell_str(rec.get("country")) or None,
            "notes": _cell_str(rec.get("notes")) or None,
        })
    # keep sheet order for rows without an explicit number
    for i, s in enumerate(stops):
        if s["position"] is None:
            s["position"] = i + 1
    stops.sort(key=lambda s: s["position"])
    for i, s in enumerate(stops, start=1):
        s["position"] = i
    return stops, errors


def parse_landmarks(file_obj):
    """-> (landmarks, errors). Each landmark carries the internal field names."""
    header, body = _read_rows(file_obj)
    if header is None:
        return [], ["El archivo está vacío."]
    mapping = _map_headers(header, LANDMARK_ALIASES)
    missing = [f for f in ("name", "architect", "city") if f not in mapping.values()]
    if missing:
        labels = {"name": "Edificio", "architect": "Arquitecto", "city": "Ciudad"}
        return [], ["Faltan columnas obligatorias: " + ", ".join(labels[m] for m in missing)
                    + ". Usa la plantilla descargada."]

    items, errors, seen = [], [], set()
    for rownum, cells in body:
        rec = {mapping[i]: cells[i] if i < len(cells) else None for i in mapping}
        if not any(_cell_str(v) for v in rec.values()):
            continue
        name = _cell_str(rec.get("name"))
        architect = _cell_str(rec.get("architect"))
        city = _cell_str(rec.get("city"))
        problems = [lbl for lbl, val in (("Edificio", name), ("Arquitecto", architect), ("Ciudad", city)) if not val]
        if problems:
            errors.append(f"fila {rownum}: falta " + " y ".join(problems))
            continue
        key = landmark_key(name, architect)
        if key in seen:
            errors.append(f"fila {rownum}: '{name}' de {architect} está repetido; se usa la primera")
            continue
        seen.add(key)
        lat, lon = _cell_float(rec.get("lat")), _cell_float(rec.get("lon"))
        if (lat is None) != (lon is None):
            errors.append(f"fila {rownum}: latitud y longitud deben ir juntas; se ignoran")
            lat = lon = None
        status = normalise(rec.get("status")) or None
        if status is not None and status not in STATUSES:
            errors.append(f"fila {rownum}: estado '{status}' no válido (posible, descartado, curado o pendiente); se ignora")
            status = None
        items.append({
            "name": name,
            "architect": architect,
            "city": city,
            "address": _cell_str(rec.get("address")) or None,
            "year": _cell_str(rec.get("year")) or None,
            "lat": lat,
            "lon": lon,
            "url_archdaily": _cell_str(rec.get("url_archdaily")) or None,
            "url_av": _cell_str(rec.get("url_av")) or None,
            "url_image1": _cell_str(rec.get("url_image1")) or None,
            "url_image2": _cell_str(rec.get("url_image2")) or None,
            "notes": _cell_str(rec.get("notes")) or None,
            "status": status,
            "name_key": key,
        })
    return items, errors
